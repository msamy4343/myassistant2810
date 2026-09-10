# ====== 🤖 مساعدي الشخصي — نسخة البحث 🔍 ======
import telebot
import requests
import json
import os
import io
import base64
import asyncio
import threading
from datetime import datetime, timedelta
import edge_tts
import pypdf
from docx import Document as DocxDocument
import openpyxl
from bs4 import BeautifulSoup

TELEGRAM_TOKEN = os.environ["TELEGRAM_TOKEN"]
GEMINI_API_KEY = os.environ["GEMINI_API_KEY"]

bot = telebot.TeleBot(TELEGRAM_TOKEN)

GEMINI_URL = "https://generativelanguage.googleapis.com/v1beta/models/gemini-3.6-flash:generateContent?key=" + GEMINI_API_KEY

TTS_VOICE = "ar-EG-ShakirNeural"

PERSONALITY = """أنت "مساعدي" — مساعد شخصي عربي ذكي.
صاحبك اسمه مصطفى، مشغول بـ: شغل + دراسة + مشروع.
قواعدك:
- ردودك ودودة ومباشرة بدون مقدمات طويلة
- لما تشرح حاجة: اشرح بالتفصيل مع أمثلة عملية
- لما يطلب تنفيذ حاجة: نفذها كاملة وجاهزة
- افتكر تفاصيل حياته واستخدمها لما تنفع
- عند الإجابة عن الملفات: استخدم معرفتك المخزنة
- لو السؤال عن معلومات حديثة (أسعار/أخبار/طقس): قول له يدور في النت بكلمة "دور" أو "ابحث"
"""

MEMORY_FILE = "memory.json"
REMINDERS_FILE = "reminders.json"
FILES_FILE = "files.json"

def load_memory():
    if os.path.exists(MEMORY_FILE):
        with open(MEMORY_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}

def save_memory():
    with open(MEMORY_FILE, "w", encoding="utf-8") as f:
        json.dump(memory, f, ensure_ascii=False, indent=2)

memory = load_memory()

# ⏰ ====== التذكيرات ======
def load_reminders():
    if os.path.exists(REMINDERS_FILE):
        with open(REMINDERS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return []

def save_reminders():
    with open(REMINDERS_FILE, "w", encoding="utf-8") as f:
        json.dump(reminders, f, ensure_ascii=False, indent=2)

reminders = load_reminders()

REMINDER_KEYWORDS = ["فكرني", "فكرنى", "ذكرني", "ذكرنى", "ذكّرني", "ذكّرنى"]

def is_reminder_request(text):
    return any(k in text for k in REMINDER_KEYWORDS)

def gemini_request(body):
    r = requests.post(GEMINI_URL, json=body)
    result = r.json()
    if "error" in result:
        error_msg = result["error"].get("message", "خطأ غير معروف")
        print("🔴 خطأ من جيميناي:", error_msg)
        raise Exception(error_msg)
    return result

def parse_reminder(text):
    now = datetime.now()
    prompt = f"""المستخدم طلب تذكير. الوقت الحالي: {now.strftime('%Y-%m-%d %H:%M')}
رسالة المستخدم: "{text}"

رد بـ JSON فقط بدون أي شرح:
{{"text": "نص التذكير المختصر بدون كلمة فكرني", "minutes": <عدد الدقايق من دلوقتي>}}

لو الوقت مش واضح خالص: {{"text": "", "minutes": -1}}"""

    body = {"contents": [{"role": "user", "parts": [{"text": prompt}]}]}
    result = gemini_request(body)
    raw = result["candidates"][0]["content"]["parts"][0]["text"].strip()
    raw = raw.replace("```json", "").replace("```", "").strip()
    try:
        data = json.loads(raw)
        return data.get("text", ""), int(data.get("minutes", -1))
    except:
        return "", -1

def fire_reminder(r):
    try:
        bot.send_message(r["chat_id"], f"⏰ تذكير: {r['text']}")
    except Exception as e:
        print("خطأ في إرسال التذكير:", e)
    reminders[:] = [x for x in reminders if x["id"] != r["id"]]
    save_reminders()

def schedule_reminder(chat_id, minutes, text):
    reminder_id = datetime.now().strftime("%Y%m%d%H%M%S%f")
    fire_at = datetime.now() + timedelta(minutes=minutes)
    reminders.append({
        "id": reminder_id,
        "chat_id": chat_id,
        "text": text,
        "fire_at": fire_at.isoformat()
    })
    save_reminders()

    timer = threading.Timer(minutes * 60, fire_reminder,
                             args=[{"id": reminder_id, "chat_id": chat_id, "text": text}])
    timer.daemon = True
    timer.start()
    return fire_at

def restore_reminders():
    now = datetime.now()
    pending = []
    for r in reminders:
        try:
            fire_time = datetime.fromisoformat(r["fire_at"])
            diff = (fire_time - now).total_seconds()
            if diff < 0:
                diff = 3
            timer = threading.Timer(diff, fire_reminder, args=[r])
            timer.daemon = True
            timer.start()
            pending.append(r)
        except Exception as e:
            print("تذكير بايظ:", e)
    reminders[:] = pending
    save_reminders()
    if pending:
        print(f"⏰ رجّعت {len(pending)} تذكير محفوظ")

# 📁 ====== معرفة الملفات ======
def load_knowledge():
    if os.path.exists(FILES_FILE):
        with open(FILES_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}

def save_knowledge():
    with open(FILES_FILE, "w", encoding="utf-8") as f:
        json.dump(knowledge, f, ensure_ascii=False, indent=2)

knowledge = load_knowledge()

def extract_text_from_file(file_bytes, file_name):
    name = file_name.lower()
    try:
        if name.endswith('.pdf'):
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(pages), len(reader.pages)
        elif name.endswith('.docx'):
            doc = DocxDocument(io.BytesIO(file_bytes))
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paras), len(paras)
        elif name.endswith('.xlsx'):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
            texts = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                texts.append(f"=== ورقة: {sheet} ===")
                for row in ws.iter_rows(values_only=True):
                    if any(c is not None for c in row):
                        texts.append(" | ".join(str(c) for c in row if c is not None))
            return "\n".join(texts), len(wb.sheetnames)
        else:
            return None, 0
    except Exception as e:
        print("🔴 خطأ في قراءة الملف:", e)
        return None, 0

def summarize_text(text):
    prompt = f"لخص النص ده في نقاط مختصرة ومنظمة بالعربي:\n\n{text[:15000]}"
    body = {"contents": [{"role": "user", "parts": [{"text": prompt}]}]}
    result = gemini_request(body)
    return result["candidates"][0]["content"]["parts"][0]["text"]

# 🔍 ====== البحث في الإنترنت ======
SEARCH_KEYWORDS = [
    "دور على", "دور لي", "دورلي", "ابحث عن", "ابحث لي", "ابحثلي",
    "بحث عن", "آخر أخبار", "اخبار", "أخبار", "سعر", "أسعار", "اسعار",
    "طقس", "الطقس", "الجو بكرة", "الجو النهاردة", "النتائج", "نتيجة"
]

def needs_search(text):
    return any(k in text for k in SEARCH_KEYWORDS)

def web_search(query, max_results=6):
    """بحث حقيقي في DuckDuckGo"""
    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"
        }
        r = requests.post("https://html.duckduckgo.com/html/",
                          data={"q": query}, headers=headers, timeout=20)
        soup = BeautifulSoup(r.text, "html.parser")
        results = []
        for res in soup.select("div.result")[:max_results]:
            title_el = res.select_one("a.result__a")
            snippet_el = res.select_one(".result__snippet")
            if title_el:
                results.append({
                    "title": title_el.get_text(strip=True),
                    "url": title_el.get("href", ""),
                    "snippet": snippet_el.get_text(strip=True) if snippet_el else ""
                })
        return results
    except Exception as e:
        print("🔴 خطأ في البحث:", e)
        return []

def search_and_answer(user_text):
    """يدور + يجيب + يلخص بالمصادر"""
    results = web_search(user_text)
    if not results:
        return "😔 معرفش ألاقي نتايج دلوقتي — جرب تاني بعد شوية أو صيغ سؤالك بشكل تاني."

    context = "\n\n".join([
        f"عنوان: {r['title']}\nملخص: {r['snippet']}\nالرابط: {r['url']}"
        for r in results
    ])

    prompt = f"""المستخدم سأل: "{user_text}"

دورت في الإنترنت ولقيت النتايج دي:

{context}

لخص الإجابة بالعربي بطريقة واضحة ومنظمة ومفيدة للمستخدم.
اذكر المصادر في الآخر (عنوان + رابط)."""

    return ask_gemini([], prompt)

# 🧠 ====== المحادثة ======
def ask_gemini(history, user_message):
    contents = []
    for msg in history:
        contents.append({"role": msg["role"], "parts": [{"text": msg["content"]}]})
    contents.append({"role": "user", "parts": [{"text": user_message}]})

    body = {
        "system_instruction": {"parts": [{"text": PERSONALITY}]},
        "contents": contents
    }
    result = gemini_request(body)
    return result["candidates"][0]["content"]["parts"][0]["text"]

def voice_to_text(ogg_bytes):
    audio_b64 = base64.b64encode(ogg_bytes).decode()
    body = {
        "contents": [{
            "parts": [
                {"text": "حوّل الرسالة الصوتية دي لنص مكتوب. اكتب النص فقط."},
                {"inline_data": {"mime_type": "audio/ogg", "data": audio_b64}}
            ]
        }]
    }
    result = gemini_request(body)
    return result["candidates"][0]["content"]["parts"][0]["text"].strip()

def text_to_speech_file(text):
    async def _make():
        communicate = edge_tts.Communicate(text, TTS_VOICE)
        await communicate.save("reply.mp3")
    asyncio.run(_make())
    audio = io.BytesIO(open("reply.mp3", "rb").read())
    audio.name = "reply.mp3"
    return audio

def process_message(user_id, user_text):
    # 📋 الملفات
    if "ملفاتي" in user_text:
        if "امسح" in user_text or "الغ" in user_text:
            knowledge.clear()
            save_knowledge()
            return "🗑️ مسحت كل معرفة الملفات."
        if not knowledge:
            return "مفيش ملفات مخزنة لسه 📁 ابعتلي أي PDF أو Word أو Excel أو صورة!"
        lines = [f"• {name} — {info['summary'][:80]}..." for name, info in knowledge.items()]
        return "📁 ملفاتك المخزنة:\n" + "\n".join(lines)

    # ⏰ التذكيرات
    if "تذكيراتي" in user_text or "التذكيرات" in user_text:
        if "امسح" in user_text or "الغ" in user_text:
            reminders.clear()
            save_reminders()
            return "🗑️ مسحت كل التذكيرات."
        if not reminders:
            return "مفيش أي تذكيرات مجدولة دلوقتي 🙂"
        lines = [f"• {r['text']} — الساعة {datetime.fromisoformat(r['fire_at']).strftime('%H:%M')}" for r in reminders]
        return "⏰ تذكيراتك المجدولة:\n" + "\n".join(lines)

    if is_reminder_request(user_text):
        text, minutes = parse_reminder(user_text)
        if minutes < 0:
            return "الوقت مش واضح 😅 قولي بالظبط إمتى — مثلاً: فكرني الساعة 8 مساءً"
        fire_at = schedule_reminder(user_id, minutes, text)
        return f"⏰ تمام! هفكرك بـ «{text}» الساعة {fire_at.strftime('%H:%M')}"

    # 🔍 البحث في الإنترنت
    if needs_search(user_text):
        return search_and_answer(user_text)

    # 💬 محادثة عادية (+ سياق الملفات)
    if user_id not in memory:
        memory[user_id] = []
    history = memory[user_id][-30:]

    message_with_context = user_text
    for name, info in knowledge.items():
        if name in user_text:
            context = f"\n\n[معرفتك من ملف «{name}»:\nالملخص: {info['summary']}\nجزء من النص:\n{info.get('text', '')[:8000]}]"
            message_with_context += context
            break

    reply = ask_gemini(history, message_with_context)

    memory[user_id].append({"role": "user", "content": user_text})
    memory[user_id].append({"role": "model", "content": reply})
    save_memory()
    return reply

@bot.message_handler(commands=['start'])
def start(message):
    bot.reply_to(message,
        "أهلاً يا مصطفى! 👋\n\n"
        "📝 اكتبلي أي حاجة\n"
        "🎤 ابعتلي فويس\n"
        "⏰ قوللي: فكرني الساعة كذا\n"
        "📁 ابعتلي ملف أو صورة\n"
        "🔍 قوللي: دور على... أو اسألني عن أخبار وأسعار\n\n"
        "/مسح — مسح الذاكرة")

@bot.message_handler(commands=['مسح'])
def reset(message):
    memory[str(message.chat.id)] = []
    save_memory()
    bot.reply_to(message, "✅ مسحت الذاكرة — بداية جديدة!")

@bot.message_handler(content_types=['document'])
def handle_document(message):
    try:
        bot.send_chat_action(message.chat.id, 'typing')
        file_name = message.document.file_name or "file"
        print(f"📥 ملف وصل: {file_name}")

        file_info = bot.get_file(message.document.file_id)
        file_bytes = bot.download_file(file_info.file_path)

        text, count = extract_text_from_file(file_bytes, file_name)
        if text is None:
            bot.reply_to(message,
                "⚠️ الصيغة دي مش مدعومة لسه.\n"
                "المدعوم: PDF، Word (docx)، Excel (xlsx) — والصور 📁")
            return

        if not text.strip():
            bot.reply_to(message,
                "🤔 الملف وصل لكن النص طلع فاضي — غالباً PDF ماسوح.\n"
                "جرّب تصور صفحاته كصور وابعتها لي 📸")
            return

        caption = message.caption or ""
        summary = summarize_text(text)

        if caption:
            question_with_text = f"بناء على ملف «{file_name}»، أجب على: {caption}\n\nنص الملف:\n{text[:20000]}"
            reply = ask_gemini([], question_with_text)
        else:
            reply = (f"✅ قريت «{file_name}» ({count} جزء) وخزنته في معرفتي.\n\n"
                     f"📋 الملخص:\n{summary}\n\n"
                     f"💡 اسألني أي وقت: «إيه أهم النقط في {file_name}؟»")

        knowledge[file_name] = {
            "summary": summary,
            "text": text[:25000],
            "count": count
        }
        save_knowledge()

        bot.reply_to(message, reply)

    except Exception as e:
        print("🔴 خطأ في الملف:", e)
        bot.reply_to(message, "⚠️ حصل خطأ في قراءة الملف")

@bot.message_handler(content_types=['photo'])
def handle_photo(message):
    try:
        bot.send_chat_action(message.chat.id, 'typing')
        file_id = message.photo[-1].file_id
        file_info = bot.get_file(file_id)
        img_bytes = bot.download_file(file_info.file_path)

        question = message.caption or "وصف الصورة دي بإيجاز، ولو فيها نص اقراه."
        img_b64 = base64.b64encode(img_bytes).decode()

        body = {
            "contents": [{
                "parts": [
                    {"text": question},
                    {"inline_data": {"mime_type": "image/jpeg", "data": img_b64}}
                ]
            }]
        }
        result = gemini_request(body)
        reply = result["candidates"][0]["content"]["parts"][0]["text"]

        bot.reply_to(message, reply)
    except Exception as e:
        print("🔴 خطأ في الصورة:", e)
        bot.reply_to(message, "⚠️ حصل خطأ في تحليل الصورة")

@bot.message_handler(func=lambda m: m.content_type == 'text')
def chat(message):
    user_id = str(message.chat.id)
    try:
        bot.send_chat_action(message.chat.id, 'typing')
        reply = process_message(user_id, message.text)
        bot.reply_to(message, reply)
    except Exception as e:
        print("خطأ:", e)
        bot.reply_to(message, "⚠️ حصل خطأ، جرب تاني")

@bot.message_handler(content_types=['voice'])
def handle_voice(message):
    user_id = str(message.chat.id)
    try:
        bot.send_chat_action(message.chat.id, 'typing')
        file_info = bot.get_file(message.voice.file_id)
        ogg_bytes = bot.download_file(file_info.file_path)

        user_text = voice_to_text(ogg_bytes)
        print("🎤 قلت:", user_text)

        reply = process_message(user_id, user_text)

        bot.send_chat_action(message.chat.id, 'record_voice')
        audio = text_to_speech_file(reply)
        bot.send_audio(message.chat.id, audio)
    except Exception as e:
        print("خطأ في الصوت:", e)
        bot.reply_to(message, "⚠️ حصل خطأ، ابعتها تاني")

# 🚀 الإقلاع
restore_reminders()
print("🤖 المساعد شغال (نص + صوت + تذكيرات + ملفات + بحث 🔍)!")
bot.infinity_polling()
