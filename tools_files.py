import io
import base64

import pypdf
from docx import Document as DocxDocument
import openpyxl

from memory_store import knowledge, save, memory
from gemini_core import gemini_request

bot = None  # يتحدد من assistant.py

def extract_text_from_file(file_bytes, file_name):
    name = file_name.lower()
    try:
        if name.endswith('.pdf'):
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(pages), len(reader.pages)
        elif name.endswith('.docx'):
            doc = DocxDocument(io.BytesIO(file_bytes))
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paras), len(paras)
        elif name.endswith('.xlsx'):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
            texts = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                texts.append(f"=== ورقة: {sheet} ===")
                for row in ws.iter_rows(values_only=True):
                    if any(c is not None for c in row):
                        texts.append(" | ".join(str(c) for c in row if c is not None))
            return "\n".join(texts), len(wb.sheetnames)
        else:
            return None, 0
    except Exception as e:
        print("🔴 خطأ في قراءة الملف:", e)
        return None, 0

def summarize_text(text):
    prompt = f"لخص النص ده في نقاط مختصرة ومنظمة بالعربي:\n\n{text[:15000]}"
    body = {"contents": [{"role": "user", "parts": [{"text": prompt}]}]}
    result = gemini_request(body)
    return result["candidates"][0]["content"]["parts"][0]["text"]

# 📄 إنشاء ملف Word
def create_word_doc(title, content):
    doc = DocxDocument()
    doc.add_heading(title, 0)
    for line in content.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("#"):
            doc.add_heading(line.lstrip("# "), 1)
        elif line.startswith(("•", "-", "*")):
            doc.add_paragraph(line.lstrip("•-* "), style="List Bullet")
        else:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

WORD_KEYWORDS = ["ملخص word", "ملف word", "ملخص وورد", "ملف وورد", "اعمل word", "اعملي word"]

def is_word_request(text):
    t = text.lower()
    return any(k in t for k in WORD_KEYWORDS)

def make_conversation_word(user_id, chat_id):
    history = memory.get(user_id, [])[-20:]
    if not history:
        return "مفيش محادثة ألخصها لسه 🙂 اتكلم معايا الأول"
    conv_text = "\n".join([("👤: " if m["role"] == "user" else "🤖: ") + m["content"] for m in history])
    summary = summarize_text("المحادثة دي بين مصطفى ومساعده:\n" + conv_text)
    buf = create_word_doc("ملخص المحادثة مع مساعدي", summary)
    buf.name = "ملخص_المحادثة.docx"
    if bot:
        bot.send_document(chat_id, buf)
    return "📄 بعتلك ملف الوورد بالملخص!"

def handle_files_text(user_id, user_text):
    if "ملفاتي" in user_text:
        if "امسح" in user_text or "الغ" in user_text:
            knowledge.clear()
            save("knowledge", knowledge)
            return "🗑️ مسحت كل معرفة الملفات."
        if not knowledge:
            return "مفيش ملفات مخزنة 📁 ابعتلي أي PDF / Word / Excel / صورة!"
        lines = [f"• {name} — {info['summary'][:80]}..." for name, info in knowledge.items()]
        return "📁 ملفاتك المخزنة:\n" + "\n".join(lines)
    return None

def handle_document(message):
    try:
        bot.send_chat_action(message.chat.id, 'typing')
        file_name = message.document.file_name or "file"
        print(f"📥 ملف وصل: {file_name}")

        file_info = bot.get_file(message.document.file_id)
        file_bytes = bot.download_file(file_info.file_path)

        text, count = extract_text_from_file(file_bytes, file_name)
        if text is None:
            bot.reply_to(message, "⚠️ الصيغة مش مدعومة.\nالمتاح: PDF / Word / Excel — والصور 📁")
            return
        if not text.strip():
            bot.reply_to(message, "🤔 النص فاضي — غالباً PDF ماسوح. صوّره وابعتله كصور 📸")
            return

        caption = message.caption or ""
        summary = summarize_text(text)

        if caption:
            prompt = f"بناء على ملف «{file_name}»، أجب على: {caption}\n\nنص الملف:\n{text[:20000]}"
            body = {"contents": [{"role": "user", "parts": [{"text": prompt}]}]}
            result = gemini_request(body)
            reply = result["candidates"][0]["content"]["parts"][0]["text"]
        else:
            reply = (f"✅ قريت «{file_name}» ({count} جزء) وخزنته.\n\n"
                     f"📋 الملخص:\n{summary}\n\n💡 اسألني: «إيه أهم النقط في {file_name}؟»")

        knowledge[file_name] = {"summary": summary, "text": text[:25000], "count": count}
        save("knowledge", knowledge)

        bot.reply_to(message, reply)
    except Exception as e:
        print("🔴 خطأ في الملف:", e)
        bot.reply_to(message, "⚠️ حصل خطأ في قراءة الملف")

def handle_photo(message):
    try:
        bot.send_chat_action(message.chat.id, 'typing')
        file_id = message.photo[-1].file_id
        file_info = bot.get_file(file_id)
        img_bytes = bot.download_file(file_info.file_path)

        question = message.caption or "وصف الصورة بإيجاز، ولو فيها نص اقراه."
        img_b64 = base64.b64encode(img_bytes).decode()

        body = {
            "contents": [{
                "parts": [
                    {"text": question},
                    {"inline_data": {"mime_type": "image/jpeg", "data": img_b64}}
                ]
            }]
        }
        result = gemini_request(body)
        reply = result["candidates"][0]["content"]["parts"][0]["text"]
        bot.reply_to(message, reply)
    except Exception as e:
        print("🔴 خطأ في الصورة:", e)
        bot.reply_to(message, "⚠️ حصل خطأ في تحليل الصورة")
